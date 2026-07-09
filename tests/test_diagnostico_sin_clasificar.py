"""Pruebas del diagnostico tecnico para registros sin clasificar."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from logica.exportador_excel import ExportadorExcel
from reglas.evaluador_condiciones import ContextoMotor, FarmaciaAtributos
from reglas.motor_reglas_avanzado import CondicionMotor, MotorReglasAvanzado, ReglaMotor


def _contexto_base(*, listas: dict[int, frozenset[str]] | None = None) -> ContextoMotor:
    return ContextoMotor(
        farmacias={
            "FARM_ORIGEN": FarmaciaAtributos("INTERNA", True, True, False),
            "FARM_DESTINO": FarmaciaAtributos("INTERNA", True, True, False),
            "FARM_OTRA": FarmaciaAtributos("EXTERNA", True, False, True),
        },
        listas=listas or {},
    )


def _clasificar(
    dataframe: pd.DataFrame,
    reglas: list[ReglaMotor],
    contexto: ContextoMotor | None = None,
) -> pd.DataFrame:
    resultado = MotorReglasAvanzado(
        proveedor_reglas=lambda: reglas,
        proveedor_contexto=lambda: contexto or _contexto_base(),
    ).clasificar(dataframe)
    return resultado.dataframe_resultado


def _fila_base(**overrides: str) -> dict[str, str]:
    fila = {
        "TIPO_TRANSACCION": "KNOWN",
        "TIPO_ORIGEN": "INVENTORY",
        "ORG_ORIGEN": "FARM_ORIGEN",
        "ORG_DESTINO": "FARM_DESTINO",
        "ARTICULO": "A1",
        "DESCRIPCION": "Articulo prueba",
        "SUBINVENTARIO": "",
    }
    fila.update(overrides)
    return fila


def test_registro_clasificado_tiene_estado_y_regla_aplicada() -> None:
    reglas = [
        ReglaMotor(
            nombre="regla_clasifica_known",
            prioridad=1,
            tipologia_resultado="AJUSTE",
            condiciones=[CondicionMotor("TIPO_TRANSACCION", "IGUAL", ["KNOWN"])],
        )
    ]

    df = _clasificar(pd.DataFrame([_fila_base()]), reglas)

    assert df.loc[0, "ESTADO_CLASIFICACION"] == "CLASIFICADO"
    assert df.loc[0, "REGLA_APLICADA"] == "regla_clasifica_known"
    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"] == ""
    assert df.loc[0, "SUGERENCIA_ACCION"] == ""


def test_registro_sin_clasificar_tiene_motivo_y_sugerencia() -> None:
    df = _clasificar(pd.DataFrame([_fila_base()]), [])

    assert df.loc[0, "TIPOLOGIA_FINAL"] == "SIN_CLASIFICAR"
    assert df.loc[0, "ESTADO_CLASIFICACION"] == "SIN_CLASIFICAR_JUSTIFICADO"
    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"]
    assert df.loc[0, "SUGERENCIA_ACCION"]


def test_diagnostico_detecta_org_origen_vacio() -> None:
    df = _clasificar(pd.DataFrame([_fila_base(ORG_ORIGEN="")]), [])

    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"] == "ORG_ORIGEN vacio."
    assert "archivo fuente" in df.loc[0, "SUGERENCIA_ACCION"]


def test_diagnostico_detecta_org_destino_vacio() -> None:
    df = _clasificar(pd.DataFrame([_fila_base(ORG_DESTINO="")]), [])

    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"] == "ORG_DESTINO vacio."
    assert "archivo fuente" in df.loc[0, "SUGERENCIA_ACCION"]


def test_diagnostico_detecta_farmacia_externa_nueva() -> None:
    df = _clasificar(pd.DataFrame([_fila_base(ORG_ORIGEN="FARM_EXTERNA_NUEVA")]), [])

    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"] == "ORG_ORIGEN parece farmacia externa nueva."
    assert "externa" in df.loc[0, "SUGERENCIA_ACCION"]


def test_diagnostico_detecta_tipo_transaccion_desconocido() -> None:
    reglas = [
        ReglaMotor(
            nombre="regla_tipo_known",
            prioridad=1,
            tipologia_resultado="AJUSTE",
            condiciones=[CondicionMotor("TIPO_TRANSACCION", "IGUAL", ["KNOWN"])],
        )
    ]

    df = _clasificar(pd.DataFrame([_fila_base(TIPO_TRANSACCION="UNKNOWN")]), reglas)

    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"] == "TIPO_TRANSACCION no esta contemplado por reglas activas."
    assert "regla" in df.loc[0, "SUGERENCIA_ACCION"]


def test_diagnostico_detecta_articulo_fuera_de_listas_configurables() -> None:
    contexto = _contexto_base(listas={7: frozenset({"A1"})})
    reglas = [
        ReglaMotor(
            nombre="regla_lista_liquidos",
            prioridad=1,
            tipologia_resultado="LIQUIDOS",
            condiciones=[
                CondicionMotor("TIPO_TRANSACCION", "IGUAL", ["KNOWN"]),
                CondicionMotor("ARTICULO", "EN_LISTA", [], lista_id=7),
            ],
        )
    ]

    df = _clasificar(pd.DataFrame([_fila_base(ARTICULO="B2")]), reglas, contexto)

    assert df.loc[0, "MOTIVO_SIN_CLASIFICAR"] == "ARTICULO no existe en listas configurables relevantes."
    assert "LIQUIDOS" in df.loc[0, "SUGERENCIA_ACCION"]


def test_diagnostico_reporta_regla_parcial_y_condiciones_fallidas() -> None:
    reglas = [
        ReglaMotor(
            nombre="salidas_prestamos_externos",
            prioridad=1,
            tipologia_resultado="SALIDA_PRESTAMO_EXTERNO",
            condiciones=[
                CondicionMotor("TIPO_TRANSACCION", "IGUAL", ["KNOWN"]),
                CondicionMotor("TIPO_ORIGEN", "IGUAL", ["INVENTORY"]),
                CondicionMotor("ORG_ORIGEN", "IGUAL", ["FARM_OTRA"]),
            ],
        )
    ]

    df = _clasificar(pd.DataFrame([_fila_base()]), reglas)

    assert df.loc[0, "REGLA_CANDIDATA"] == "salidas_prestamos_externos"
    assert "ORG_ORIGEN" in df.loc[0, "CONDICIONES_FALLIDAS"]
    assert df.loc[0, "PORCENTAJE_COINCIDENCIA_REGLA"] > 0
    assert "Coincidencia parcial" in df.loc[0, "MOTIVO_SIN_CLASIFICAR"]


def test_exportacion_sin_clasificar_incluye_diagnostico_obligatorio() -> None:
    sin_clasificar = pd.DataFrame(
        [
            {
                "TIPO_TRANSACCION": "UNKNOWN",
                "TIPO_ORIGEN": "INVENTORY",
                "ORG_ORIGEN": "",
                "ORG_DESTINO": "FARM_DESTINO",
                "ARTICULO": "A1",
                "DESCRIPCION": "Articulo prueba",
                "TIPOLOGIA_PRELIMINAR": "SIN_CLASIFICAR",
                "TIPOLOGIA_FINAL": "SIN_CLASIFICAR",
            }
        ]
    )

    with tempfile.TemporaryDirectory(prefix="farmaflow_test_", dir=Path.cwd()) as carpeta:
        resultado = ExportadorExcel().exportar(
            dataframe_original=sin_clasificar,
            detalle_clasificado=sin_clasificar,
            resumen_tipologia=pd.DataFrame({"TIPOLOGIA": ["SIN_CLASIFICAR"], "CONTEO": [1]}),
            resumen_farmacia=pd.DataFrame({"FARMACIA": ["FARM_DESTINO"], "CONTEO": [1]}),
            sin_clasificar=sin_clasificar,
            ruta_salida=carpeta,
            nombre_base_archivo="diagnostico",
        )

        assert resultado.exito
        libro = load_workbook(resultado.ruta_salida, read_only=True)
        hoja = libro["SIN_CLASIFICAR"]
        encabezados = [celda.value for celda in next(hoja.iter_rows(min_row=1, max_row=1))]
        fila = [celda.value for celda in next(hoja.iter_rows(min_row=2, max_row=2))]
        valores = dict(zip(encabezados, fila, strict=False))

        assert "MOTIVO_SIN_CLASIFICAR" in encabezados
        assert "SUGERENCIA_ACCION" in encabezados
        assert "REGLA_CANDIDATA" in encabezados
        assert "CONDICIONES_FALLIDAS" in encabezados
        assert valores["MOTIVO_SIN_CLASIFICAR"] == "ORG_ORIGEN vacio."
        assert valores["SUGERENCIA_ACCION"]


def test_no_quedan_sin_clasificar_sin_motivo() -> None:
    df = _clasificar(
        pd.DataFrame(
            [
                _fila_base(),
                _fila_base(ORG_ORIGEN=""),
                _fila_base(TIPO_TRANSACCION="UNKNOWN"),
            ]
        ),
        [],
    )
    sin = df["TIPOLOGIA_FINAL"].eq("SIN_CLASIFICAR")

    assert not df.loc[sin, "MOTIVO_SIN_CLASIFICAR"].fillna("").astype(str).eq("").any()
    assert not df.loc[sin, "SUGERENCIA_ACCION"].fillna("").astype(str).eq("").any()
