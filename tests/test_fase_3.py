"""Pruebas funcionales minimas para la Fase 3 de exportacion operativa."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from logica.exportador_excel import ExportadorExcel
from logica.generador_resumen import (
    generar_articulos_unicos_por_lista,
    generar_cruce_farmacia_tipologia,
    generar_detalle_filtrado_por_lista,
    generar_estructuras_exportacion,
    generar_resumen_farmacia,
    generar_resumen_tipologia,
    generar_sin_clasificar,
)
from utilidades.rutas import asegurar_directorio, construir_nombre_archivo_salida


class PruebasFaseTres(unittest.TestCase):
    """Cobertura minima para resumenes y exportacion Excel."""

    def setUp(self) -> None:
        self.dataframe_original = pd.DataFrame(
            [
                {
                    "ARTICULO": "ART_001",
                    "DESCRIPCION": "Jeringa",
                    "ORG DESTINO": "Farmacia Norte",
                    "TIPO TRANSACCION": "RMA Receipt",
                    "MOTIVO": "",
                },
                {
                    "ARTICULO": "ART_002",
                    "DESCRIPCION": "Solucion salina",
                    "ORG DESTINO": "Farmacia Norte",
                    "TIPO TRANSACCION": "Sales Order Issue",
                    "MOTIVO": "",
                },
                {
                    "ARTICULO": "ART_003",
                    "DESCRIPCION": "Equipo infusion",
                    "ORG DESTINO": "Farmacia Sur",
                    "TIPO TRANSACCION": "Ajuste",
                    "MOTIVO": "Revision",
                },
            ]
        )
        self.dataframe_procesado = pd.DataFrame(
            [
                {
                    "FARMACIA_DETECTADA": "FARMACIA_NORTE",
                    "TIPOLOGIA_PRELIMINAR": "DEVOLUCIONES",
                    "REGLA_APLICADA": "regla_devolucion",
                },
                {
                    "FARMACIA_DETECTADA": "FARMACIA_NORTE",
                    "TIPOLOGIA_PRELIMINAR": "DISPENSACION",
                    "REGLA_APLICADA": "regla_dispensacion",
                },
                {
                    "FARMACIA_DETECTADA": "FARMACIA_SUR",
                    "TIPOLOGIA_PRELIMINAR": "SIN_CLASIFICAR",
                    "REGLA_APLICADA": "",
                },
            ]
        )
        self.estructuras = generar_estructuras_exportacion(
            dataframe_original=self.dataframe_original,
            dataframe_procesado=self.dataframe_procesado,
        )

    def test_generacion_resumen_tipologia(self) -> None:
        resumen = generar_resumen_tipologia(self.estructuras["detalle_clasificado"])

        cantidades = dict(zip(resumen["TIPOLOGIA_PRELIMINAR"], resumen["CANTIDAD"], strict=False))

        self.assertEqual(cantidades["DEVOLUCIONES"], 1)
        self.assertEqual(cantidades["DISPENSACION"], 1)
        self.assertEqual(cantidades["SIN_CLASIFICAR"], 1)

    def test_generacion_resumen_farmacia(self) -> None:
        resumen = generar_resumen_farmacia(self.estructuras["detalle_clasificado"])

        cantidades = dict(zip(resumen["FARMACIA_DETECTADA"], resumen["CANTIDAD"], strict=False))

        self.assertEqual(cantidades["FARMACIA_NORTE"], 2)
        self.assertEqual(cantidades["FARMACIA_SUR"], 1)

    def test_generacion_cruce_tipologia_farmacia(self) -> None:
        cruce = generar_cruce_farmacia_tipologia(self.estructuras["detalle_clasificado"])

        fila_devoluciones = cruce.loc[cruce["TIPOLOGIA_PRELIMINAR"] == "DEVOLUCIONES"].iloc[0]
        fila_sin_clasificar = cruce.loc[cruce["TIPOLOGIA_PRELIMINAR"] == "SIN_CLASIFICAR"].iloc[0]

        self.assertEqual(int(fila_devoluciones["FARMACIA_NORTE"]), 1)
        self.assertEqual(int(fila_sin_clasificar["FARMACIA_SUR"]), 1)

    def test_filtrado_por_lista_de_articulos(self) -> None:
        filtrado = generar_detalle_filtrado_por_lista(
            self.estructuras["detalle_clasificado"],
            ["ART_001", "ART 003"],
        )

        self.assertEqual(filtrado["ARTICULO"].tolist(), ["ART_001", "ART_003"])

    def test_listado_unico_de_articulos_por_lista(self) -> None:
        detalle = pd.concat(
            [
                self.estructuras["detalle_clasificado"],
                self.estructuras["detalle_clasificado"].iloc[[0]],
            ],
            ignore_index=True,
        )

        listado = generar_articulos_unicos_por_lista(detalle, ["ART_001", "ART 003"])

        self.assertEqual(list(listado.columns), ["CODIGO", "DESCRIPCION", "CONTEO"])
        self.assertEqual(listado.to_dict("records"), [
            {"CODIGO": "ART_001", "DESCRIPCION": "Jeringa", "CONTEO": 2},
            {"CODIGO": "ART_003", "DESCRIPCION": "Equipo infusion", "CONTEO": 1},
        ])

    def test_filtrado_de_sin_clasificar(self) -> None:
        sin_clasificar = generar_sin_clasificar(self.estructuras["detalle_clasificado"])

        self.assertEqual(len(sin_clasificar), 1)
        self.assertEqual(sin_clasificar.iloc[0]["FARMACIA_DETECTADA"], "FARMACIA_SUR")
        self.assertEqual(sin_clasificar.iloc[0]["TIPOLOGIA_PRELIMINAR"], "SIN_CLASIFICAR")

    def test_validacion_del_nombre_de_archivo_generado(self) -> None:
        nombre_archivo = construir_nombre_archivo_salida("reporte abril: farmacia?.xlsx")

        self.assertRegex(nombre_archivo, r"^clasificado_reporte_abril_farmacia_\d{8}_\d{6}\.xlsx$")

    def test_creacion_automatica_del_directorio_de_salida(self) -> None:
        with tempfile.TemporaryDirectory() as carpeta_temporal:
            ruta_directorio = Path(carpeta_temporal) / "subcarpeta" / "salidas"

            self.assertFalse(ruta_directorio.exists())
            resultado = asegurar_directorio(ruta_directorio)

            self.assertTrue(ruta_directorio.exists())
            self.assertEqual(resultado, ruta_directorio)

    def test_exportacion_exitosa_con_hojas_obligatorias(self) -> None:
        exportador = ExportadorExcel()

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            ruta_salida = Path(carpeta_temporal) / "salidas_generadas"
            resultado = exportador.exportar(
                dataframe_original=self.dataframe_original,
                detalle_clasificado=self.estructuras["detalle_clasificado"],
                resumen_tipologia=self.estructuras["resumen_tipologia"],
                resumen_farmacia=self.estructuras["resumen_farmacia"],
                cruce_tipologia_farmacia=self.estructuras["cruce_tipologia_farmacia"],
                sin_clasificar=self.estructuras["sin_clasificar"],
                ruta_salida=ruta_salida,
                nombre_base_archivo="movimientos abril.xlsx",
            )

            self.assertTrue(resultado.exito)
            self.assertTrue(Path(resultado.ruta_salida).exists())
            self.assertEqual(resultado.hojas_generadas, list(ExportadorExcel.HOJAS_OBLIGATORIAS))

            libro = load_workbook(resultado.ruta_salida)
            self.assertEqual(libro.sheetnames, list(ExportadorExcel.HOJAS_OBLIGATORIAS))
            self.assertEqual(libro["DETALLE_CLASIFICADO"].freeze_panes, "A2")
            self.assertTrue(bool(libro["RESUMEN_TIPOLOGIA"].auto_filter.ref))
            self.assertEqual(libro["SIN_CLASIFICAR"].max_row, 2)

    def test_exportacion_agrega_hojas_articulos_si_tienen_datos(self) -> None:
        exportador = ExportadorExcel()
        liquidos = pd.concat(
            [
                self.estructuras["detalle_clasificado"].iloc[[0]],
                self.estructuras["detalle_clasificado"].iloc[[0]],
            ],
            ignore_index=True,
        )

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            resultado = exportador.exportar(
                dataframe_original=self.dataframe_original,
                detalle_clasificado=self.estructuras["detalle_clasificado"],
                resumen_tipologia=self.estructuras["resumen_tipologia"],
                resumen_farmacia=self.estructuras["resumen_farmacia"],
                cruce_tipologia_farmacia=self.estructuras["cruce_tipologia_farmacia"],
                sin_clasificar=self.estructuras["sin_clasificar"],
                liquidos=liquidos,
                mce_cirugia=self.estructuras["detalle_clasificado"].iloc[[1]],
                ruta_salida=Path(carpeta_temporal),
                nombre_base_archivo="movimientos abril.xlsx",
            )

            libro = load_workbook(resultado.ruta_salida)

            self.assertIn("Liquidos Cirugia", libro.sheetnames)
            self.assertIn("MCE_CIRUGIA", libro.sheetnames)
            self.assertEqual(
                [celda.value for celda in libro["Liquidos Cirugia"][1]],
                ["CODIGO", "DESCRIPCION", "CONTEO"],
            )
            self.assertEqual(
                [celda.value for celda in libro["Liquidos Cirugia"][2]],
                ["ART_001", "Jeringa", 2],
            )
            self.assertEqual(
                [celda.value for celda in libro["MCE_CIRUGIA"][2]],
                ["ART_002", "Solucion salina", 1],
            )

    def test_exportador_no_muta_dataframe_original(self) -> None:
        """Garantia de que el exportador (sin deep copy) no altera la entrada."""
        exportador = ExportadorExcel()
        original_id = id(self.dataframe_original)
        original_columnas = list(self.dataframe_original.columns)
        original_filas = len(self.dataframe_original)

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            exportador.exportar(
                dataframe_original=self.dataframe_original,
                detalle_clasificado=self.estructuras["detalle_clasificado"],
                resumen_tipologia=self.estructuras["resumen_tipologia"],
                resumen_farmacia=self.estructuras["resumen_farmacia"],
                cruce_tipologia_farmacia=self.estructuras["cruce_tipologia_farmacia"],
                sin_clasificar=self.estructuras["sin_clasificar"],
                ruta_salida=Path(carpeta_temporal),
                nombre_base_archivo="test",
            )

        # El dataframe original debe seguir siendo la misma instancia con las
        # mismas columnas y misma cantidad de filas.
        self.assertEqual(id(self.dataframe_original), original_id)
        self.assertEqual(list(self.dataframe_original.columns), original_columnas)
        self.assertEqual(len(self.dataframe_original), original_filas)

    def test_exportador_emite_progreso(self) -> None:
        """Verifica que el callback de progreso se invoca y avanza a 100%."""
        exportador = ExportadorExcel()
        progreso_recibido: list[tuple[int, str]] = []

        def callback(porcentaje: int, mensaje: str) -> None:
            progreso_recibido.append((porcentaje, mensaje))

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            resultado = exportador.exportar(
                dataframe_original=self.dataframe_original,
                detalle_clasificado=self.estructuras["detalle_clasificado"],
                resumen_tipologia=self.estructuras["resumen_tipologia"],
                resumen_farmacia=self.estructuras["resumen_farmacia"],
                cruce_tipologia_farmacia=self.estructuras["cruce_tipologia_farmacia"],
                sin_clasificar=self.estructuras["sin_clasificar"],
                ruta_salida=Path(carpeta_temporal),
                nombre_base_archivo="test",
                progreso_callback=callback,
            )

        self.assertTrue(resultado.exito)
        self.assertGreater(len(progreso_recibido), 1)
        # Debe llegar al 100% al final.
        self.assertEqual(progreso_recibido[-1][0], 100)

    def test_exportador_respeta_cancelacion(self) -> None:
        """Si el callback de cancelacion devuelve True, la exportacion se aborta."""
        exportador = ExportadorExcel()

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            resultado = exportador.exportar(
                dataframe_original=self.dataframe_original,
                detalle_clasificado=self.estructuras["detalle_clasificado"],
                resumen_tipologia=self.estructuras["resumen_tipologia"],
                resumen_farmacia=self.estructuras["resumen_farmacia"],
                cruce_tipologia_farmacia=self.estructuras["cruce_tipologia_farmacia"],
                sin_clasificar=self.estructuras["sin_clasificar"],
                ruta_salida=Path(carpeta_temporal),
                nombre_base_archivo="test_cancelado",
                cancelado_callback=lambda: True,
            )

        self.assertFalse(resultado.exito)
        self.assertIn("cancelada", resultado.mensaje.lower())

    def test_generar_detalle_no_muta_original_y_no_deep_copy(self) -> None:
        """generar_detalle_clasificado no debe alterar el dataframe original."""
        from logica.generador_resumen import generar_detalle_clasificado

        df_original = pd.DataFrame({"COL_A": [1, 2, 3], "COL_B": ["x", "y", "z"]})
        df_procesado = pd.DataFrame({
            "COL_A": [1, 2, 3],
            "COL_B": ["x", "y", "z"],
            "FARMACIA_DETECTADA": ["F1", "F2", "F1"],
            "TIPOLOGIA_PRELIMINAR": ["T1", "T2", "SIN_CLASIFICAR"],
            "REGLA_APLICADA": ["r1", "r2", ""],
        })
        original_columnas = list(df_original.columns)

        detalle = generar_detalle_clasificado(df_original, df_procesado)

        # Original intacto
        self.assertEqual(list(df_original.columns), original_columnas)
        # Detalle agrega las 3 columnas auxiliares
        for columna in ("FARMACIA_DETECTADA", "TIPOLOGIA_PRELIMINAR", "REGLA_APLICADA"):
            self.assertIn(columna, detalle.columns)


if __name__ == "__main__":
    unittest.main()
