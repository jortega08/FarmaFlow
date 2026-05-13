"""Genera un Excel de prueba con movimientos sinteticos para validar el clasificador.

Cubre las 15 tipologias del cuadro de filtros + casos SIN_CLASIFICAR + conteo ciclico.
Salida: ``datos_prueba/movimientos_simulados.xlsx`` (hoja ``Movimientos``).
"""

from __future__ import annotations

import random
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

RAIZ_PROYECTO = Path(__file__).resolve().parent.parent
RUTA_SALIDA = RAIZ_PROYECTO / "datos_prueba" / "movimientos_simulados.xlsx"

COLUMNAS = [
    "ARTICULO", "DESCRIPCION", "SUBINVENTARIO", "ORG_ORIGEN", "ORG_DESTINO",
    "FECHA_TRANSACCION", "TIPO_TRANSACCION", "CANT_TRANSACCION", "UDM_TRANSACCION",
    "TIPO_ORIGEN", "ORIGEN", "MOTIVO", "NUMERO_ENC_TRANSAC", "AJUSTE", "PEDIDO", "REFERENCIA",
]

ARTICULOS = [
    (19991, "AGUA ESTERIL IRRIGACION BOLSA X 3000 ML"),
    (19942, "LACTATO RINGER USO HOSPITALARIO X 1000 ML"),
    (38839, "CLORURO DE SODIO 0.9% BOLSA X 500 ML"),
    (45211, "DEXTROSA 5% BOLSA X 500 ML"),
    (52301, "ACETAMINOFEN 500 MG TABLETA"),
    (60127, "AMOXICILINA 500 MG CAPSULA"),
    (71045, "IBUPROFENO 400 MG TABLETA"),
    (84502, "OMEPRAZOL 20 MG CAPSULA"),
    (91108, "LOSARTAN 50 MG TABLETA"),
    (10283, "METFORMINA 850 MG TABLETA"),
]

ORGS_INTERNAS = [
    "16_FARMA_FARMACIA_INTERNA_CRS",
    "256_FARMA_URGENCIAS_CRS",
    "47_FARMA_ALMACEN_CIRUGIA_CRS",
]
ORGS_CEDI = [
    "1058_CRUZ_VERDE_CEDI_COTA_ALMACEN_PRINCIPAL",
    "337_FARMA_NO_VENDIBLE",
    "901_CRUZ_VERDE_ALMACEN_DE_AVERIADOS",
    "991_CRUZ_VERDE_PICKING_DEVOLUCIONES",
]
ORGS_EXTERNAS = [
    "20_FARMA_CLINICA_COUNTRY",
    "27_FARMA_CLINICA_VERSALLES",
    "26_FARMA_FARMACIA_INTERNA_CSB",
    "228_FARMA_FARMACIA_INTERNA_CUC",
    "315_CRUZ_VERDE_CANCEROLOGICO",
    "368_FARMACIA_CLINICENTRO_CHIA",
    "489_FARMACIA_INTERNA_CLINICA_PEDIATRICA",
    "920_CRUZ_VERDE_COMPENSAR_CARRERA_15",
]
ORGS_CEDI_TRASLADOS = [
    "373_CRUZ_VERDE_ALMACEN_DEVOLUCIONES",
    "800_CRUZ_VERDE_RETIRO",
]
ORG_CONSIGNACION = "102_FARMA_CONSIGN_PROV_CRS"
ORG_CENTRAL_PREP = "104_FARMA_CENTRAL_PREP_BOG"
ORG_CENTRAL_RERE = "702_CRUZ_VERDE_CENTRAL_REEMPAQUE_REENVASE"

UDMS = ["UND", "ML", "TBT", "CAP", "BLT"]
SUBINVENTARIOS = ["DISPONIBLE", "AVE_VENCID", "TRANSITO"]


def fecha_aleatoria(rng: random.Random) -> datetime:
    base = datetime(2026, 1, 1)
    return base + timedelta(days=rng.randint(0, 90), hours=rng.randint(0, 23))


def encabezado(rng: random.Random, prefijo: str) -> str:
    return f"{prefijo}-{rng.randint(1000, 9999)}"


def fila_base(rng: random.Random) -> dict:
    articulo, descripcion = rng.choice(ARTICULOS)
    return {
        "ARTICULO": articulo,
        "DESCRIPCION": descripcion,
        "SUBINVENTARIO": "DISPONIBLE",
        "ORG_ORIGEN": "",
        "ORG_DESTINO": "",
        "FECHA_TRANSACCION": fecha_aleatoria(rng),
        "TIPO_TRANSACCION": "",
        "CANT_TRANSACCION": rng.randint(1, 60),
        "UDM_TRANSACCION": rng.choice(UDMS),
        "TIPO_ORIGEN": "",
        "ORIGEN": "",
        "MOTIVO": "",
        "NUMERO_ENC_TRANSAC": encabezado(rng, "ENC"),
        "AJUSTE": "",
        "PEDIDO": "",
        "REFERENCIA": "",
    }


def generar_filas(rng: random.Random) -> list[dict]:
    filas: list[dict] = []

    def agregar(cantidad: int, **campos):
        for _ in range(cantidad):
            fila = fila_base(rng)
            fila.update(campos)
            filas.append(fila)

    # 1. DISPENSACION_AL_PACIENTE
    for destino in ORGS_INTERNAS:
        agregar(2,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="SALES_ORDER_ISSUE",
                TIPO_ORIGEN="SALES_ORDER",
                ORIGEN="MASIVO_CONSUMO.ORDER",
                PEDIDO=f"SO-{rng.randint(100000, 999999)}")

    # 2. DEVOLUCIONES
    for destino in ORGS_INTERNAS:
        agregar(2,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="RMA_RECEIPT",
                TIPO_ORIGEN="RMA",
                REFERENCIA=f"RMA-{rng.randint(1000, 9999)}")

    # 3. ENTRADAS_PRESTAMOS_INTERNOS (origen y destino son farmacias internas)
    for _ in range(6):
        origen = rng.choice(ORGS_INTERNAS)
        destino = rng.choice(ORGS_INTERNAS)
        agregar(1,
                ORG_ORIGEN=origen,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="INTRANSIT_RECEIPT",
                TIPO_ORIGEN="INVENTORY",
                NUMERO_ENC_TRANSAC=encabezado(rng, "TRX-IR"))

    # 4. SALIDAS_PRESTAMOS_INTERNOS
    for _ in range(6):
        origen = rng.choice(ORGS_INTERNAS)
        destino = rng.choice(ORGS_INTERNAS)
        agregar(1,
                ORG_ORIGEN=origen,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="INTRANSIT_SHIPMENT",
                TIPO_ORIGEN="INVENTORY",
                NUMERO_ENC_TRANSAC=encabezado(rng, "TRX-IS"))

    # 5. ENTRADAS_CEDI
    for origen in ORGS_CEDI:
        agregar(2,
                ORG_ORIGEN=origen,
                ORG_DESTINO=rng.choice(["16_FARMA_FARMACIA_INTERNA_CRS", "47_FARMA_ALMACEN_CIRUGIA_CRS"]),
                TIPO_TRANSACCION="ACCOUNT_ALIAS_RECEIPT",
                TIPO_ORIGEN="ACCOUNT_ALIAS")

    # 6. ORDENES_DE_COMPRA
    for destino in ["16_FARMA_FARMACIA_INTERNA_CRS", "47_FARMA_ALMACEN_CIRUGIA_CRS"]:
        agregar(2,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="PO_RECEIPT",
                TIPO_ORIGEN="PURCHASE_ORDER",
                REFERENCIA=f"PO-{rng.randint(10000, 99999)}")

    # 7. SALIDAS_PRESTAMOS_EXTERNOS (dos patrones)
    for _ in range(8):
        origen = rng.choice(ORGS_EXTERNAS)
        destino = rng.choice(ORGS_INTERNAS)
        agregar(1,
                ORG_ORIGEN=origen,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="INTRANSIT_SHIPMENT",
                TIPO_ORIGEN="INVENTORY")
    for _ in range(4):
        origen = rng.choice(ORGS_EXTERNAS)
        destino = rng.choice(ORGS_INTERNAS)
        agregar(1,
                ORG_ORIGEN=origen,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="INT_ORDER_INTR_SHIP",
                TIPO_ORIGEN="INTERNAL_ORDER")

    # 8. ENTRADAS_CONSIGNACION
    for destino in ORGS_INTERNAS:
        agregar(1,
                ORG_ORIGEN=ORG_CONSIGNACION,
                ORG_DESTINO=destino,
                TIPO_TRANSACCION="INTRANSIT_RECEIPT",
                TIPO_ORIGEN="INVENTORY")

    # 9. ENTRADAS_CENTRAL_PREP (4 combinaciones)
    combos_prep = [
        ("INT_REQ_INTR_RCPT", "INTERNAL_REQUISITION"),
        ("INTRANSIT_RECEIPT", "INTERNAL_REQUISITION"),
        ("INT_REQ_INTR_RCPT", "INVENTORY"),
        ("INTRANSIT_RECEIPT", "INVENTORY"),
    ]
    for tipo_trx, tipo_origen in combos_prep:
        agregar(1,
                ORG_ORIGEN=ORG_CENTRAL_PREP,
                ORG_DESTINO="16_FARMA_FARMACIA_INTERNA_CRS",
                TIPO_TRANSACCION=tipo_trx,
                TIPO_ORIGEN=tipo_origen)

    # 10. ENTRADAS_CENTRAL_RERE
    agregar(2,
            ORG_ORIGEN=ORG_CENTRAL_RERE,
            ORG_DESTINO="16_FARMA_FARMACIA_INTERNA_CRS",
            TIPO_TRANSACCION="INTRANSIT_RECEIPT",
            TIPO_ORIGEN="INVENTORY")

    # 11. CONTEO_CICLICO (fallback)
    agregar(3,
            ORG_DESTINO="16_FARMA_FARMACIA_INTERNA_CRS",
            TIPO_TRANSACCION="CYCLE_COUNT_ADJUST",
            TIPO_ORIGEN="CYCLE_COUNT",
            AJUSTE="SI")

    # 12. PRESTAMOS_BOPOS
    for tipo_trx in ["ACCOUNT_ALIAS_RECEIPT", "ACCOUNT_ALIAS_ISSUE"]:
        agregar(2,
                ORG_DESTINO="16_FARMA_FARMACIA_INTERNA_CRS",
                TIPO_TRANSACCION=tipo_trx,
                TIPO_ORIGEN="ACCOUNT_ALIAS",
                ORIGEN="MOVIMIENTOS BOPOS-EBS")

    # 13. AJUSTES_BASE_CAPTURA
    for origen_aa in ["AA PRODUCTOS AVERIADOS DE PUNTOS",
                      "AA PROXIMOS A VENCER DE PUNTOS",
                      "ANEXO 6 SALIDA DESTRUCCION MCE"]:
        agregar(2,
                ORG_DESTINO=rng.choice(["16_FARMA_FARMACIA_INTERNA_CRS", "47_FARMA_ALMACEN_CIRUGIA_CRS"]),
                TIPO_TRANSACCION="ACCOUNT_ALIAS_ISSUE",
                TIPO_ORIGEN="ACCOUNT_ALIAS",
                ORIGEN=origen_aa,
                MOTIVO=rng.choice(["08_VENCIDO", "09_AVERIADO"]),
                AJUSTE="SI")

    # 14. AJUSTES_INCONSISTENCIAS (issue + receipt)
    for tipo_trx in ["ACCOUNT_ALIAS_ISSUE", "ACCOUNT_ALIAS_RECEIPT"]:
        for origen_inc in ["LEGALIZACION INCONSISTENCIAS DESPACHO",
                           "LEGALIZACION INCONSIST DEVOLUCIONES PTO"]:
            agregar(1,
                    ORG_DESTINO=rng.choice(["16_FARMA_FARMACIA_INTERNA_CRS", "47_FARMA_ALMACEN_CIRUGIA_CRS"]),
                    TIPO_TRANSACCION=tipo_trx,
                    TIPO_ORIGEN="ACCOUNT_ALIAS",
                    ORIGEN=origen_inc)

    # 15. TRASLADOS_HACIA_CEDI
    for origen in ORGS_CEDI_TRASLADOS:
        agregar(2,
                ORG_ORIGEN=origen,
                ORG_DESTINO=rng.choice(["16_FARMA_FARMACIA_INTERNA_CRS", "47_FARMA_ALMACEN_CIRUGIA_CRS"]),
                TIPO_TRANSACCION="INTRANSIT_SHIPMENT",
                TIPO_ORIGEN="INVENTORY")

    # 16. SIN_CLASIFICAR (combinaciones que no encajan)
    agregar(3,
            ORG_ORIGEN="999_FARMA_INEXISTENTE",
            ORG_DESTINO="888_FARMA_DESCONOCIDA",
            TIPO_TRANSACCION="MISC_ISSUE",
            TIPO_ORIGEN="ADJUSTMENT")
    agregar(2,
            ORG_DESTINO="16_FARMA_FARMACIA_INTERNA_CRS",
            TIPO_TRANSACCION="ACCOUNT_ALIAS_ISSUE",
            TIPO_ORIGEN="ACCOUNT_ALIAS",
            ORIGEN="ORIGEN_NO_CONFIGURADO")

    rng.shuffle(filas)
    return filas


def escribir_excel(filas: list[dict], ruta: Path) -> None:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Movimientos"
    hoja.append(COLUMNAS)
    for fila in filas:
        hoja.append([fila[col] for col in COLUMNAS])
    for indice, columna in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = max(14, len(columna) + 2)
    hoja.freeze_panes = "A2"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


def main() -> int:
    rng = random.Random(20260428)
    filas = generar_filas(rng)
    escribir_excel(filas, RUTA_SALIDA)
    print(f"Generadas {len(filas)} filas en {RUTA_SALIDA.relative_to(RAIZ_PROYECTO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
