"""Genera ``configuracion/reglas_tipologia.json`` a partir del cuadro Excel de criterios.

Uso:
    python scripts/generar_reglas_desde_excel.py "<ruta del xlsx>"

Si no se pasa ruta, busca ``CRITERIOS DE FILTROS .xlsx`` en ``Downloads`` del usuario.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl

RAIZ_PROYECTO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ_PROYECTO))

from utilidades.texto import normalizar_texto  # noqa: E402

RUTA_SALIDA = RAIZ_PROYECTO / "configuracion" / "reglas_tipologia.json"

ENCABEZADOS_COL = [
    "ARTICULO", "DESCRIPCION", "SUBINVENTARIO", "ORG_ORIGEN", "ORG_DESTINO",
    "FECHA_TRANSACCION", "ORG_PROPIETARIA", "ID_TRANSACCION", "CANT_TRANSACCION",
    "UDM_TRANSACCION", "UDM_PRIMARIO", "CANT_PRIMARIA", "TIPOLOGIA",
    "TIPO_TRANSACCION", "TIPO_ORIGEN", "ORIGEN", "MOTIVO", "NUMERO_ENVIO",
    "ID_TRANSACCION_TRANS", "AJUSTE", "PEDIDO", "REFERENCIA",
]
IDX = {nombre: pos for pos, nombre in enumerate(ENCABEZADOS_COL)}

CAMPOS_CONDICION = [
    "SUBINVENTARIO", "ORG_ORIGEN", "ORG_DESTINO",
    "TIPO_TRANSACCION", "TIPO_ORIGEN", "ORIGEN", "MOTIVO",
]

ETIQUETA_A_CODIGO = {
    "SALES ORDER ISSUE": "SALES_ORDER_ISSUE",
    "SALES ORDER": "SALES_ORDER",
    "RMA RECEIPT": "RMA_RECEIPT",
    "RMA": "RMA",
    "INTRANSIT RECEIPT": "INTRANSIT_RECEIPT",
    "INTRANSIT SHIPMENT": "INTRANSIT_SHIPMENT",
    "INVENTORY": "INVENTORY",
    "ACCOUNT ALIAS RECEIPT": "ACCOUNT_ALIAS_RECEIPT",
    "ACCOUNT ALIAS ISSUE": "ACCOUNT_ALIAS_ISSUE",
    "ACCOUNT ALIAS": "ACCOUNT_ALIAS",
    "PO RECEIPT": "PO_RECEIPT",
    "PURCHASE ORDER": "PURCHASE_ORDER",
    "INT ORDER INTR SHIP": "INT_ORDER_INTR_SHIP",
    "INTERNAL ORDER": "INTERNAL_ORDER",
    "INT REQ INTR RCPT": "INT_REQ_INTR_RCPT",
    "INTERNAL REQUISITION": "INTERNAL_REQUISITION",
}

CAMPOS_TRADUCIBLES = {"TIPO_TRANSACCION", "TIPO_ORIGEN"}

CORRECCIONES_TYPOS = {
    "ANEXO 6 SALIDA DESTRUCCIEN MCE": "ANEXO 6 SALIDA DESTRUCCION MCE",
}

PRIORIDADES = {
    "AJUSTES_BASE_CAPTURA": 10,
    "AJUSTES_INCONSISTENCIAS": 10,
    "PRESTAMOS_BOPOS": 10,
    "ENTRADAS_CEDI": 20,
    "ENTRADAS_CONSIGNACION": 20,
    "ENTRADAS_CENTRAL_PREP": 20,
    "ENTRADAS_CENTRAL_RERE": 20,
    "TRASLADOS_HACIA_CEDI": 20,
    "ENTRADAS_PRESTAMOS_INTERNOS": 30,
    "SALIDAS_PRESTAMOS_INTERNOS": 30,
    "SALIDAS_PRESTAMOS_EXTERNOS": 40,
    "DISPENSACION_AL_PACIENTE": 50,
    "DEVOLUCIONES": 50,
    "ORDENES_DE_COMPRA": 50,
}


def codificar(campo: str, valor: str) -> str:
    """Traduce etiquetas display a código si aplica al campo."""
    valor_norm = normalizar_texto(valor)
    valor_norm = CORRECCIONES_TYPOS.get(valor_norm, valor_norm)
    if campo in CAMPOS_TRADUCIBLES:
        return ETIQUETA_A_CODIGO.get(valor_norm, valor_norm)
    return valor_norm


def detectar_secciones(filas: list[tuple]) -> list[tuple[int, int, str]]:
    """Identifica los bloques de tipología devolviendo (inicio_datos, fin_datos, nombre)."""
    encabezados: list[tuple[int, str]] = []
    for indice, fila in enumerate(filas):
        if (
            fila[0]
            and fila[1] is None
            and fila[2] is None
            and fila[IDX["TIPOLOGIA"]] is None
            and fila[IDX["TIPO_TRANSACCION"]] is None
            and str(fila[0]).strip().upper() != "ARTICULO"
        ):
            encabezados.append((indice, str(fila[0]).strip()))

    encabezados = [(idx, nombre) for idx, nombre in encabezados if "CUADRO DE REGLA" not in nombre.upper()]

    secciones: list[tuple[int, int, str]] = []
    for posicion, (indice, nombre) in enumerate(encabezados):
        fin = encabezados[posicion + 1][0] if posicion + 1 < len(encabezados) else len(filas)
        secciones.append((indice + 1, fin, nombre))
    return secciones


def filas_de_seccion(filas: list[tuple], inicio: int, fin: int) -> list[tuple]:
    """Devuelve solo las filas con condiciones (descarta encabezados ARTICULO y vacíos)."""
    filtradas = []
    for fila in filas[inicio:fin]:
        if not fila or all(c is None or str(c).strip() == "" for c in fila):
            continue
        if str(fila[0] or "").strip().upper() == "ARTICULO":
            continue
        filtradas.append(fila)
    return filtradas


def construir_condiciones(filas_seccion: list[tuple]) -> list[dict[str, list[str]]]:
    """Agrupa filas que comparten ``TIPO_TRANSACCION`` + ``TIPO_ORIGEN`` y unifica el resto."""
    grupos: dict[tuple[str, str], dict[str, set[str]]] = {}
    for fila in filas_seccion:
        clave = (
            codificar("TIPO_TRANSACCION", fila[IDX["TIPO_TRANSACCION"]]),
            codificar("TIPO_ORIGEN", fila[IDX["TIPO_ORIGEN"]]),
        )
        condiciones = grupos.setdefault(clave, {campo: set() for campo in CAMPOS_CONDICION})
        for campo in CAMPOS_CONDICION:
            valor = codificar(campo, fila[IDX[campo]])
            if valor:
                condiciones[campo].add(valor)

    resultados = []
    for (tipo_trx, tipo_origen), valores in grupos.items():
        condiciones_finales: dict[str, list[str]] = {}
        if tipo_trx:
            condiciones_finales["TIPO_TRANSACCION"] = [tipo_trx]
        if tipo_origen:
            condiciones_finales["TIPO_ORIGEN"] = [tipo_origen]
        for campo in CAMPOS_CONDICION:
            if campo in {"TIPO_TRANSACCION", "TIPO_ORIGEN"}:
                continue
            if valores[campo]:
                condiciones_finales[campo] = sorted(valores[campo])
        resultados.append(condiciones_finales)
    return resultados


def nombre_a_tipologia(nombre_seccion: str) -> str:
    return normalizar_texto(nombre_seccion).replace(" ", "_").replace("-", "_")


def construir_reglas(ruta_excel: Path) -> list[dict[str, Any]]:
    libro = openpyxl.load_workbook(ruta_excel, data_only=True)
    hoja = libro[libro.sheetnames[0]]
    filas = list(hoja.iter_rows(values_only=True))
    secciones = detectar_secciones(filas)

    reglas: list[dict[str, Any]] = []
    for inicio, fin, nombre_seccion in secciones:
        filas_seccion = filas_de_seccion(filas, inicio, fin)
        if not filas_seccion:
            continue
        tipologia = nombre_a_tipologia(nombre_seccion)
        if tipologia == "CONTEOS_CICLICOS":
            continue  # se conserva el fallback existente más abajo
        condiciones_lista = construir_condiciones(filas_seccion)
        prioridad = PRIORIDADES.get(_alias_tipologia(tipologia), 100)
        sufijo = "" if len(condiciones_lista) == 1 else None
        for posicion, condiciones in enumerate(condiciones_lista, start=1):
            if not condiciones:
                continue
            nombre_regla = tipologia.lower()
            if sufijo is None:
                nombre_regla = f"{nombre_regla}_{posicion}"
            reglas.append({
                "nombre_regla": nombre_regla,
                "prioridad": prioridad,
                "condiciones": condiciones,
                "resultado": _alias_tipologia(tipologia),
            })

    reglas.append({
        "nombre_regla": "conteo_ciclico_fallback",
        "prioridad": 999,
        "condiciones": {
            "TIPO_TRANSACCION": ["CYCLE_COUNT_ADJUST"],
            "TIPO_ORIGEN": ["CYCLE_COUNT"],
        },
        "resultado": "CONTEO_CICLICO",
    })

    reglas.sort(key=lambda regla: (regla["prioridad"], regla["nombre_regla"]))
    return reglas


def _alias_tipologia(tipologia: str) -> str:
    """Ajusta nombres de sección a la convención usada en los resultados."""
    tabla = {
        "DISPENSACION": "DISPENSACION_AL_PACIENTE",
    }
    return tabla.get(tipologia, tipologia)


def main() -> int:
    if len(sys.argv) >= 2:
        ruta_excel = Path(sys.argv[1])
    else:
        ruta_excel = Path.home() / "Downloads" / "CRITERIOS DE FILTROS .xlsx"
    if not ruta_excel.exists():
        print(f"No se encontró el Excel en: {ruta_excel}", file=sys.stderr)
        return 1

    reglas = construir_reglas(ruta_excel)
    RUTA_SALIDA.write_text(
        json.dumps(reglas, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Generadas {len(reglas)} reglas en {RUTA_SALIDA.relative_to(RAIZ_PROYECTO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
