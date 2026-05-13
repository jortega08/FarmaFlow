"""Exportacion operativa de resultados a Excel multihoja.

Optimizado para archivos grandes (100k+ filas):
- No clona dataframes (pandas.to_excel no muta los datos de entrada).
- El autosize de columnas usa un muestreo (no recorre todas las filas).
- Acepta callbacks de progreso y cancelacion para integrarse con un worker Qt.
"""

from __future__ import annotations

import logging
from collections.abc import Callable
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from logica.generador_resumen import generar_resumen_articulos
from modelos.resultado_exportacion import ResultadoExportacion
from utilidades.mensajes import MensajesInterfaz
from utilidades.rutas import asegurar_directorio, construir_nombre_archivo_salida


# Limite de filas a samplear para calcular ancho de columna automatico.
# Iterar todas las celdas en archivos de 100k filas era el cuello de botella.
_MAX_FILAS_AUTOSIZE = 1000


class ExportadorExcel:
    """Genera un archivo Excel multihoja con salida operativa."""

    HOJAS_OBLIGATORIAS = (
        "ORIGINAL",
        "DETALLE_CLASIFICADO",
        "RESUMEN_TIPOLOGIA",
        "RESUMEN_FARMACIA",
        "CRUCE_TIPOLOGIA_FARMACIA",
        "SIN_CLASIFICAR",
    )

    ANCHO_MINIMO = 12
    ANCHO_MAXIMO = 40

    def __init__(self) -> None:
        self._logger = logging.getLogger(__name__)
        self._relleno_encabezado = PatternFill(fill_type="solid", fgColor="1F4E78")
        self._fuente_encabezado = Font(color="FFFFFF", bold=True)
        self._alineacion_encabezado = Alignment(horizontal="center", vertical="center")
        borde = Side(style="thin", color="D9E2F3")
        self._borde_encabezado = Border(left=borde, right=borde, top=borde, bottom=borde)

    def exportar(
        self,
        dataframe_original: pd.DataFrame,
        detalle_clasificado: pd.DataFrame,
        resumen_tipologia: pd.DataFrame,
        resumen_farmacia: pd.DataFrame,
        cruce_tipologia_farmacia: pd.DataFrame | None = None,
        sin_clasificar: pd.DataFrame | None = None,
        ruta_salida: Path | str = ".",
        nombre_base_archivo: str = "",
        hojas_seleccionadas: set[str] | None = None,
        cruce_farmacia_tipologia: pd.DataFrame | None = None,
        liquidos: pd.DataFrame | None = None,
        mce_cirugia: pd.DataFrame | None = None,
        mcs_cirugia: pd.DataFrame | None = None,
        progreso_callback: Callable[[int, str], None] | None = None,
        cancelado_callback: Callable[[], bool] | None = None,
    ) -> ResultadoExportacion:
        """Exporta los dataframes en un archivo Excel con formato basico.

        ``progreso_callback`` recibe (porcentaje 0-100, mensaje). ``cancelado_callback``
        debe devolver True si la exportacion fue cancelada por el usuario.
        """
        if detalle_clasificado is None or detalle_clasificado.empty:
            mensaje = MensajesInterfaz.ERROR_SIN_DATOS_EXPORTAR
            self._logger.warning("Se intento exportar sin detalle clasificado.")
            return ResultadoExportacion(exito=False, mensaje=mensaje)

        try:
            directorio_salida = asegurar_directorio(ruta_salida)
        except Exception as error:  # noqa: BLE001
            mensaje = f"{MensajesInterfaz.ERROR_CARPETA_SALIDA}: {error}"
            self._logger.exception("No fue posible preparar el directorio de salida: %s", ruta_salida)
            return ResultadoExportacion(exito=False, mensaje=mensaje)

        nombre_archivo = construir_nombre_archivo_salida(nombre_base=nombre_base_archivo)
        ruta_archivo = directorio_salida / nombre_archivo

        # IMPORTANTE: usar referencias directas, no deep copies.
        # pandas.to_excel no muta el dataframe de entrada y mantener 6 copies de
        # 100k filas cada una hace explotar la memoria.
        cruce_exportar = cruce_tipologia_farmacia if cruce_tipologia_farmacia is not None else cruce_farmacia_tipologia
        if cruce_exportar is None:
            cruce_exportar = pd.DataFrame()
        if sin_clasificar is None:
            sin_clasificar = pd.DataFrame()

        hojas_a_exportar: dict[str, pd.DataFrame] = {
            "ORIGINAL": dataframe_original,
            "DETALLE_CLASIFICADO": detalle_clasificado,
            "RESUMEN_TIPOLOGIA": resumen_tipologia,
            "RESUMEN_FARMACIA": resumen_farmacia,
            "CRUCE_TIPOLOGIA_FARMACIA": cruce_exportar,
            "SIN_CLASIFICAR": sin_clasificar,
        }
        if liquidos is not None and not liquidos.empty:
            hojas_a_exportar["Liquidos Cirugia"] = _preparar_hoja_articulos(liquidos)
        hoja_mce = mce_cirugia if mce_cirugia is not None else mcs_cirugia
        if hoja_mce is not None and not hoja_mce.empty:
            hojas_a_exportar["MCE_CIRUGIA"] = _preparar_hoja_articulos(hoja_mce)
        if hojas_seleccionadas is not None:
            hojas_seleccionadas = set(hojas_seleccionadas)
            if "LIQUIDOS" in hojas_seleccionadas:
                hojas_seleccionadas.add("Liquidos Cirugia")
            if "LIQUIDOS_CIRUGIA" in hojas_seleccionadas:
                hojas_seleccionadas.add("Liquidos Cirugia")
            hojas_a_exportar = {
                nombre: dataframe
                for nombre, dataframe in hojas_a_exportar.items()
                if nombre in hojas_seleccionadas
            }
            if not hojas_a_exportar:
                mensaje = MensajesInterfaz.ERROR_SIN_DATOS_EXPORTAR
                self._logger.warning("Se intento exportar sin hojas seleccionadas.")
                return ResultadoExportacion(exito=False, mensaje=mensaje)

        total_hojas = len(hojas_a_exportar)
        self._logger.info("Iniciando exportacion Excel en: %s (%s hojas)", ruta_archivo, total_hojas)
        self._emitir_progreso(progreso_callback, 0, "Preparando archivo Excel...")

        # Comprobacion temprana antes de abrir el escritor: cancelar dentro del
        # ``with pd.ExcelWriter`` rompe el __exit__ porque openpyxl exige al
        # menos una hoja visible al guardar.
        if self._fue_cancelado(cancelado_callback):
            return ResultadoExportacion(
                exito=False,
                mensaje="Exportacion cancelada por el usuario.",
            )

        cancelado = False
        try:
            with pd.ExcelWriter(ruta_archivo, engine="openpyxl") as escritor:
                for indice, (nombre_hoja, dataframe) in enumerate(hojas_a_exportar.items()):
                    porcentaje_inicio = int(indice * 100 / total_hojas)
                    self._emitir_progreso(
                        progreso_callback,
                        porcentaje_inicio,
                        f"Generando hoja {nombre_hoja} ({len(dataframe):,} filas)...",
                    )
                    dataframe.to_excel(escritor, sheet_name=nombre_hoja, index=False)
                    self._aplicar_formato_hoja(escritor.book[nombre_hoja])
                    if self._fue_cancelado(cancelado_callback):
                        cancelado = True
                        break
        except Exception as error:  # noqa: BLE001
            mensaje = f"{MensajesInterfaz.ERROR_EXPORTACION}: {error}"
            self._logger.exception("Fallo la exportacion del archivo Excel: %s", ruta_archivo)
            return ResultadoExportacion(exito=False, mensaje=mensaje)

        if cancelado:
            try:
                ruta_archivo.unlink(missing_ok=True)
            except OSError:
                pass
            return ResultadoExportacion(
                exito=False,
                mensaje="Exportacion cancelada por el usuario.",
            )

        self._emitir_progreso(progreso_callback, 100, "Archivo generado correctamente.")

        hojas_generadas = list(hojas_a_exportar.keys())
        self._logger.info(
            "Exportacion finalizada. Hojas generadas: %s. Ruta final: %s",
            ", ".join(hojas_generadas),
            ruta_archivo,
        )

        return ResultadoExportacion(
            exito=True,
            mensaje=f"{MensajesInterfaz.EXPORTACION_EXITOSA} {ruta_archivo}",
            ruta_salida=str(ruta_archivo),
            nombre_archivo=nombre_archivo,
            hojas_generadas=hojas_generadas,
            cantidad_registros_exportados=int(len(detalle_clasificado)),
        )

    # ------------------------------------------------------------------
    # Formato y autosize
    # ------------------------------------------------------------------

    def _aplicar_formato_hoja(self, hoja: Worksheet) -> None:
        """Aplica formato sobrio y legible a una hoja ya generada.

        Para hojas grandes evita iterar todas las celdas: el calculo de ancho
        de columna se hace por muestreo de las primeras ``_MAX_FILAS_AUTOSIZE``
        filas (encabezado y muestra de datos).
        """
        if hoja.max_row < 1 or hoja.max_column < 1:
            return

        hoja.freeze_panes = "A2"

        # Encabezado: estilo siempre.
        for celda in hoja[1]:
            celda.fill = self._relleno_encabezado
            celda.font = self._fuente_encabezado
            celda.alignment = self._alineacion_encabezado
            celda.border = self._borde_encabezado

        if hoja.max_row >= 1:
            hoja.auto_filter.ref = hoja.dimensions

        # Autosize: usar muestreo en hojas grandes para no recorrer 100k filas.
        filas_a_muestrear = min(hoja.max_row, _MAX_FILAS_AUTOSIZE)
        for columna_indice in range(1, hoja.max_column + 1):
            letra_columna = hoja.cell(row=1, column=columna_indice).column_letter
            longitud_maxima = 0
            for fila_indice in range(1, filas_a_muestrear + 1):
                valor = hoja.cell(row=fila_indice, column=columna_indice).value
                if valor is not None:
                    longitud = len(str(valor))
                    if longitud > longitud_maxima:
                        longitud_maxima = longitud
            ancho_ajustado = min(
                max(longitud_maxima + 2, self.ANCHO_MINIMO), self.ANCHO_MAXIMO
            )
            hoja.column_dimensions[letra_columna].width = ancho_ajustado

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _emitir_progreso(
        callback: Callable[[int, str], None] | None,
        porcentaje: int,
        mensaje: str,
    ) -> None:
        if callback is None:
            return
        try:
            callback(porcentaje, mensaje)
        except Exception:  # noqa: BLE001
            # No dejar que un error en el callback rompa la exportacion.
            pass

    @staticmethod
    def _fue_cancelado(callback: Callable[[], bool] | None) -> bool:
        if callback is None:
            return False
        try:
            return bool(callback())
        except Exception:  # noqa: BLE001
            return False


def _preparar_hoja_articulos(dataframe: pd.DataFrame) -> pd.DataFrame:
    columnas_resumen = ["CODIGO", "DESCRIPCION", "CONTEO"]
    if all(columna in dataframe.columns for columna in columnas_resumen):
        return dataframe.loc[:, columnas_resumen].copy()
    return generar_resumen_articulos(dataframe)
